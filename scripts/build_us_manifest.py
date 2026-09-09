#!/usr/bin/env python3
"""Build an Amazon US ASIN manifest from the authorized workbook export."""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path
from urllib.parse import urlsplit

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - exercised in environments without openpyxl
    load_workbook = None
    _OPENPYXL_ERROR = exc

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = ROOT / "lingxing_asin_links.xlsx"
DEFAULT_OUTPUT = ROOT / "amazon_us_asin_manifest.csv"
HEADERS = ["asin", "url", "marketplace", "source_site_label", "source_workbook"]
ASIN_RE = re.compile(r"^[A-Z0-9]{10}$")


def _find_header_row(ws):
    for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = {str(value).strip() for value in row if value is not None}
        if {"ASIN", "Amazon 商品链接", "站点"}.issubset(values):
            indices = {str(value).strip(): index for index, value in enumerate(row) if value is not None}
            return row_number, indices["ASIN"], indices["Amazon 商品链接"], indices["站点"]
    return None


def _amazon_us_asin(url: str) -> str | None:
    parts = urlsplit(url.strip())
    if parts.scheme.lower() != "https" or parts.hostname not in {"amazon.com", "www.amazon.com"}:
        return None
    match = re.search(r"/(?:dp|gp/product|product)/([A-Za-z0-9]{10})(?:/|$)", parts.path, re.IGNORECASE)
    if not match:
        return None
    asin = match.group(1).upper()
    return asin if ASIN_RE.fullmatch(asin) else None


def build_manifest(input_path: Path = DEFAULT_INPUT, output_path: Path = DEFAULT_OUTPUT, force: bool = False) -> int:
    if output_path.exists() and not force:
        raise FileExistsError(f"输出已存在，使用 --force 才能覆盖: {output_path}")
    if load_workbook is None:
        raise RuntimeError(f"需要 openpyxl 才能读取 XLSX: {_OPENPYXL_ERROR}")
    if not input_path.exists():
        raise FileNotFoundError(f"输入工作簿不存在: {input_path}")

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    records: dict[tuple[str, str], dict[str, str]] = {}
    try:
        for ws in workbook.worksheets:
            header = _find_header_row(ws)
            if header is None:
                continue
            row_number, asin_column, url_column, site_column = header
            for row in ws.iter_rows(min_row=row_number + 1, values_only=True):
                raw_url = str(row[url_column]).strip() if len(row) > url_column and row[url_column] else ""
                asin = _amazon_us_asin(raw_url)
                if asin is None:
                    continue
                source_label = str(row[site_column]).strip() if len(row) > site_column and row[site_column] else ""
                parts = urlsplit(raw_url)
                canonical_url = f"https://www.amazon.com/dp/{asin}"
                key = ("US", asin)
                records.setdefault(
                    key,
                    {
                        "asin": asin,
                        "url": canonical_url,
                        "marketplace": "US",
                        "source_site_label": source_label,
                        "source_workbook": input_path.name,
                    },
                )
    finally:
        workbook.close()

    rows = sorted(records.values(), key=lambda item: (item["marketplace"], item["asin"]))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.name}.tmp")
    with temporary.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)
    temporary.replace(output_path)
    print(f"已生成 {len(rows)} 条 US manifest: {output_path}")
    return len(rows)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--force", action="store_true", help="允许覆盖现有输出")
    args = parser.parse_args(argv)
    try:
        build_manifest(args.input, args.output, args.force)
    except (FileExistsError, FileNotFoundError, RuntimeError, OSError, ValueError) as exc:
        print(f"错误: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
