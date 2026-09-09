#!/usr/bin/env python3
"""Create the immutable owned-US manifest from the approved Sheet2 ASIN range."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path


def create_manifest(source: Path, output: Path, *, sheet: str = "Sheet2", first_row: int = 2, last_row: int = 1094) -> dict[str, object]:
    if output.exists():
        raise FileExistsError(f"refusing to overwrite manifest: {output}")
    if first_row < 1 or last_row < first_row:
        raise ValueError("invalid workbook row range")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the authorized workbook") from exc
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet]
        asins = [str(worksheet.cell(row, 1).value or "").strip().upper() for row in range(first_row, last_row + 1)]
    finally:
        workbook.close()
    if len(asins) != 1093 or any(len(asin) != 10 or not asin.isalnum() for asin in asins) or len(set(asins)) != len(asins):
        raise ValueError("authorized range must contain exactly 1,093 unique 10-character ASINs")
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("x", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["asin", "url", "marketplace", "source_site_label", "source_workbook"])
        writer.writeheader()
        for asin in asins:
            writer.writerow({"asin": asin, "url": f"https://www.amazon.com/dp/{asin}", "marketplace": "US", "source_site_label": "owned_us_warehouse", "source_workbook": source.name})
    metadata = output.with_suffix(output.suffix + ".metadata.json")
    metadata.write_text(json.dumps({
        "schema_version": "amazon-us-owned-manifest-v1",
        "created_at": datetime.now(timezone.utc).isoformat(),
        "source_path": str(source), "source_sha256": digest, "sheet": sheet,
        "range": f"A{first_row}:A{last_row}", "count": len(asins),
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"count": len(asins), "source_sha256": digest, "manifest": str(output), "metadata": str(metadata)}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        result = create_manifest(args.source, args.output)
    except (OSError, RuntimeError, ValueError) as exc:
        print(f"manifest creation failed: {exc}")
        return 2
    print(f"manifest_count={result['count']}; source_sha256={result['source_sha256']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
